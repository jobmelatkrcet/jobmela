import urllib.request
import urllib.parse
import json
import io
import openpyxl

BASE_URL = "http://127.0.0.1:8000/api"

def make_req(path, method="GET", data=None, token=None, files=None):
    url = f"{BASE_URL}{path}"
    headers = {}
    body = None

    if token:
        headers["Authorization"] = f"Token {token}"

    if data is not None and files is None:
        headers["Content-Type"] = "application/json"
        body = json.dumps(data).encode("utf-8")

    req = urllib.request.Request(url, data=body, headers=headers, method=method)

    try:
        with urllib.request.urlopen(req) as resp:
            content_type = resp.headers.get("Content-Type", "")
            content = resp.read()
            if "application/json" in content_type:
                return resp.status, json.loads(content.decode("utf-8")), resp.headers
            return resp.status, content, resp.headers
    except urllib.error.HTTPError as e:
        content = e.read().decode("utf-8")
        try:
            return e.code, json.loads(content), e.headers
        except Exception:
            return e.code, content, e.headers

print("==================================================")
print("TKRCET JOB MELA 2026 — LIVE END-TO-END HTTP TESTS")
print("==================================================")

# 1. Test Public API Root
status, data, _ = make_req("/")
assert status == 200, f"Root check failed: {status}"
print(f"✓ Root Health Check: {data['event']} ({data['status']})")

# 2. Test Public Companies List
status, data, _ = make_req("/companies/")
assert status == 200
print(f"✓ Public Companies: Found {data['count']} participating companies")
all_companies = data["results"]
assert len(all_companies) >= 3

# 3. Register Student: Naveen Kumar
test_student_email = "naveen.kumar@tkrcet.ac.in"
# Remove if exists to ensure pristine test run
import os, django
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
django.setup()
from accounts.models import User
User.objects.filter(email=test_student_email).delete()

reg_payload = {
    "full_name": "Naveen Kumar",
    "email": test_student_email,
    "mobile": "9848012345",
    "qualification": "B.Tech (CSE / IT / Allied)",
    "college": "TKR College of Engineering & Technology",
    "password": "Password@123",
    "confirm_password": "Password@123"
}
status, data, _ = make_req("/auth/register/", method="POST", data=reg_payload)
assert status == 201, f"Registration failed: {data}"
student_token = data["token"]
student_user = data["user"]
print(f"✓ Student Registered: {student_user['full_name']} ({student_user['email']})")

# 4. Duplicate Registration Prevention
status, data, _ = make_req("/auth/register/", method="POST", data=reg_payload)
assert status == 400
print(f"✓ Duplicate Email Prevention: Correctly rejected duplicate registration")

# 5. Student Login
login_payload = {"email": test_student_email, "password": "Password@123"}
status, data, _ = make_req("/auth/login/", method="POST", data=login_payload)
assert status == 200
print(f"✓ Student Login: Verified successfully")

# 6. Apply to 3 Companies (Company 1, Company 2, Company 3)
c1 = all_companies[0]["id"]
c2 = all_companies[1]["id"]
c3 = all_companies[2]["id"]

for cid in [c1, c2, c3]:
    status, data, _ = make_req("/applications/", method="POST", data={"company_id": cid}, token=student_token)
    assert status == 201, f"Failed applying to company {cid}: {data}"
print(f"✓ Multiple Company Applications: Successfully applied to 3 companies")

# 7. Duplicate Application Rejection
status, data, _ = make_req("/applications/", method="POST", data={"company_id": c1}, token=student_token)
assert status == 400
print(f"✓ Duplicate Application Prevention: Successfully rejected repeat application to same company")

# 8. Verify My Applications
status, data, _ = make_req("/applications/my/", token=student_token)
assert status == 200
assert data["total_applied"] == 3
print(f"✓ My Applications: Verified {data['total_applied']} applied companies in student profile")

# 9. Security Test: Student attempting to access Admin API (must get 403)
status, data, _ = make_req("/admin/dashboard/", token=student_token)
assert status == 403, f"Expected 403 Forbidden for student on admin endpoint, got {status}"
print(f"✓ Security Authorization: Student access to admin endpoint returned 403 Forbidden")

# 10. Admin Login
status, data, _ = make_req("/auth/login/", method="POST", data={"email": "admin@tkrcet.ac.in", "password": "admin123"})
assert status == 200
admin_token = data["token"]
print(f"✓ Admin Login: Authenticated as organizer")

# 11. Admin Dashboard Stats
status, data, _ = make_req("/admin/dashboard/", token=admin_token)
assert status == 200
print(f"✓ Admin Dashboard Metrics: Total Students={data['total_students']}, Companies={data['total_companies']}, Applications={data['total_applications']}")

# 12. Admin Company Registered Students List
status, data, _ = make_req(f"/admin/companies/{c1}/students/", token=admin_token)
assert status == 200
company_name = data["company"]["name"]
assert len(data["results"]) >= 1
print(f"✓ Company Registered Students View: {len(data['results'])} students listed for {company_name}")

# 13. Admin Company-wise Excel Export
status, content, headers = make_req(f"/admin/companies/{c1}/export/", token=admin_token)
assert status == 200
assert "openxmlformats" in headers.get("Content-Type", "")
# Read openpyxl workbook from content
wb = openpyxl.load_workbook(io.BytesIO(content))
ws = wb.active
assert ws.title == "Registered Students"
assert "TKR COLLEGE" in ws["A1"].value or "JOB MELA" in ws["A1"].value
print(f"✓ Company-wise Excel Export: Generated valid .xlsx spreadsheet ({len(content)} bytes) for {company_name}")

print("\n==================================================")
print("🎉 ALL FULL-STACK E2E HTTP INTEGRATION FLOWS PASSED!")
print("==================================================")
